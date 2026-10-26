"""Branded PDF and spreadsheet exports; all content comes from supplied records."""
from io import BytesIO
from pathlib import Path
from html import escape
from datetime import datetime,timezone
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,KeepTogether
from reportlab.graphics.shapes import Drawing,Rect,String,Line
from ess.reports.export import csv_bytes,NOTICE
from ess.risk.engine import component_summary

FONT=Path(__file__).resolve().parents[1]/'assets/fonts'
pdfmetrics.registerFont(TTFont('ESS',str(FONT/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('ESS-Bold',str(FONT/'DejaVuSans-Bold.ttf')))

def safe_frame(data):
    d=data.copy()
    for c in d:
        if pd.api.types.is_object_dtype(d[c]) or pd.api.types.is_string_dtype(d[c]):
            d[c]=d[c].map(lambda x:"'"+x if isinstance(x,str) and x.lstrip().startswith(('=','+','-','@')) else x)
    return d

def excel_bytes(sheets):
    out=BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        for i,(name,data) in enumerate(sheets.items()):
            name=''.join(c for c in name if c not in '[]:*?/\\')[:28] or f'Sheet{i+1}'
            safe_frame(data).to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]; ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font,PatternFill
            for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='17243A')
            for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(45,max(16,len(str(col[0].value))+3))
    return out.getvalue()

def report_pdf(title,metadata,tables,screening=None):
    styles=getSampleStyleSheet()
    for name in ['Normal','BodyText','Title','Heading1','Heading2']:
        styles[name].fontName='ESS-Bold' if name in ['Title','Heading1','Heading2'] else 'ESS'
    styles['Heading2'].keepWithNext=True
    styles.add(ParagraphStyle('TableHeading',parent=styles['Heading2'],keepWithNext=False))
    styles['BodyText'].fontSize=9;styles['BodyText'].leading=14
    styles['Title'].fontSize=23;styles['Title'].textColor=colors.HexColor('#17243A')
    cell=ParagraphStyle('cell',fontName='ESS',fontSize=7.5,leading=11)
    p=lambda text,style='BodyText':Paragraph(escape(str(text)),styles[style])
    story=[p('ESS SENTINEL 2.0','Heading2'),p(title,'Title'),p(datetime.now(timezone.utc).isoformat()),Spacer(1,12)]
    for k,v in metadata.items(): story.append(p(f'{k}: {v}'))
    story += [Spacer(1,12),p(NOTICE),Spacer(1,16)]
    if screening is not None and len(screening):
        counts=component_summary(screening).status.value_counts()
        chart=Drawing(470,140); palette={'NORMAL':'#37C991','WATCH':'#F5B84B','HIGH RISK':'#EE9754','CRITICAL':'#F16C76'}
        maximum=max(counts.max(),1)
        for i,level in enumerate(palette):
            y=108-i*30; val=int(counts.get(level,0))
            chart.add(String(0,y,level,fontName='ESS',fontSize=9))
            chart.add(Rect(95,y-3,280*val/maximum,15,fillColor=colors.HexColor(palette[level]),strokeColor=None))
            chart.add(String(390,y,str(val),fontName='ESS-Bold',fontSize=9))
        story += [p('Component disposition','Heading2'),chart]
        if screening.component_id.nunique()==1:
            for _,row in screening.iterrows():
                section_start=len(story)
                story.append(p(f'{row.parameter} ({row.unit}) - {row.status}','Heading2'))
                story.append(p(row.explanation))
                # Exact vector trajectory; solid observed, dashed future actual, predicted diamond-equivalent square.
                pairs=[(t,row[f'value_{t}h']) for t in [0,24,96,168] if f'value_{t}h' in row and pd.notna(row[f'value_{t}h'])]
                if pairs:
                    vals=[v for _,v in pairs]+[row.spec_min,row.spec_max]
                    if pd.notna(row.predicted_168h): vals.append(row.predicted_168h)
                    lo=min(vals); hi=max(vals); span=max(hi-lo,1)
                    xy=lambda t,v:(35+t/168*385,25+(v-lo)/span*120)
                    plot=Drawing(470,165)
                    for (ta,va),(tb,vb) in zip(pairs,pairs[1:]):
                        xa,ya=xy(ta,va); xb,yb=xy(tb,vb)
                        plot.add(Line(xa,ya,xb,yb,strokeColor=colors.HexColor('#087da0'),strokeWidth=1.6,strokeDashArray=[] if tb<=row.stage else [3,3]))
                    for t,v in pairs:
                        x,y=xy(t,v);plot.add(Rect(x-2,y-2,4,4,fillColor=colors.HexColor('#087da0'),strokeColor=None));plot.add(String(x-8,8,f'{t}h',fontName='ESS',fontSize=8))
                    for value,label,col in [(row.spec_max,'Specification','#c7464b'),(row.safety_threshold,'Safety','#a17410')]:
                        _,y=xy(0,value);plot.add(Line(30,y,430,y,strokeColor=colors.HexColor(col),strokeDashArray=[4,3]));plot.add(String(32,y+3,f'{label} {value:.3g}',fontName='ESS',fontSize=8))
                    if pd.notna(row.predicted_168h):
                        x,y=xy(168,row.predicted_168h);plot.add(Rect(x-4,y-4,8,8,fillColor=colors.HexColor('#d67d36'),strokeColor=None))
                    story += [plot,p('Solid: observed to screening stage. Dashed: retrospective future actual. Orange square: 168h prediction.')]
                section=story[section_start:];del story[section_start:];story.append(KeepTogether(section))
    for title,data in tables.items():
        story += [p(title,'TableHeading')]
        if data.empty: story.append(p('No records available for this section.'));continue
        if len(data.columns)>7: data=data.iloc[:,:7]; story.append(p('PDF shows the first seven fields; use CSV/Excel for all columns.'))
        # Full reports paginate rather than silently truncate data.
        cells=[[Paragraph(escape(str(c).replace('_',' ')),cell) for c in data.columns]]
        for row in data.itertuples(index=False,name=None):
            cells.append([Paragraph(escape('Unavailable' if pd.isna(v) else f'{v:.5g}' if isinstance(v,float) else str(v)),cell) for v in row])
        table=Table(cells,colWidths=[470/len(data.columns)]*len(data.columns),repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dce8ef')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#bfd0dd')),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        story += [table,Spacer(1,14)]
    story += [KeepTogether([p('Methodology and limitations','Heading2'),p('Robust peer statistics, early-feature Isolation Forest and lot-disjoint regression support transparent review rules. Scores are not failure probabilities. Calibration intervals are approximate under lot dependence. Synthetic evidence is not production validation. Review states do not alter engineering measurements or decisions.')])]
    buf=BytesIO()
    def footer(canvas,doc):
        canvas.setFont('ESS',8);canvas.setFillColor(colors.HexColor('#66768D'));canvas.drawString(60,28,'ESS SENTINEL / Local engineering report');canvas.drawRightString(535,28,f'{doc.page}')
    SimpleDocTemplate(buf,leftMargin=60,rightMargin=60,topMargin=45,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
    return buf.getvalue()
